from datetime import timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from pypdf import PdfWriter

from etica_ia_ers.domain import FormatoDocumento
from etica_ia_ers.ingestion import (
    DocumentoIlegibleError,
    DocumentoVacioError,
    FormatoNoSoportadoError,
    cargar_documento,
    cargar_documento_desde_ruta,
)


def _crear_pdf_con_texto(texto: str) -> bytes:
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("ascii")
    objetos[3] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(flujo), flujo)

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for numero, objeto in enumerate(objetos, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{numero} 0 obj\n".encode("ascii"))
        pdf.extend(objeto)
        pdf.extend(b"\nendobj\n")
    inicio_xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objetos) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
            f"startxref\n{inicio_xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(pdf)


def _crear_docx() -> bytes:
    documento = Document()
    documento.add_paragraph("Requisito de transparencia")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text = "Responsable"
    tabla.cell(0, 1).text = "Equipo de IA"
    documento.add_paragraph("Fin del documento")
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()


def test_carga_txt_con_metadatos_y_trazabilidad() -> None:
    contenido = "El sistema deberá explicar sus decisiones.".encode("utf-8")

    documento = cargar_documento(
        "carpeta/ers.txt", contenido, {"dominio": "crediticio"}
    )

    assert isinstance(documento.id, UUID)
    assert documento.nombre_original == "ers.txt"
    assert documento.formato_origen is FormatoDocumento.TXT
    assert documento.contenido_extraido == contenido.decode("utf-8")
    assert documento.hash_contenido == sha256(contenido).hexdigest()
    assert documento.fecha_carga.tzinfo is timezone.utc
    assert documento.metadatos == {"dominio": "crediticio"}


def test_hash_es_estable_para_el_mismo_archivo() -> None:
    contenido = b"El sistema debera registrar cada decision."

    primero = cargar_documento("ers.txt", contenido)
    segundo = cargar_documento("ers.txt", contenido)

    assert primero.id != segundo.id
    assert primero.hash_contenido == segundo.hash_contenido


def test_acepta_txt_con_bom_utf8() -> None:
    documento = cargar_documento("ers.txt", b"\xef\xbb\xbfRequisito")

    assert documento.contenido_extraido == "Requisito"


def test_extrae_paginas_pdf_en_orden() -> None:
    escritor = PdfWriter()
    escritor.append(BytesIO(_crear_pdf_con_texto("Primera pagina")))
    escritor.append(BytesIO(_crear_pdf_con_texto("Segunda pagina")))
    salida = BytesIO()
    escritor.write(salida)

    documento = cargar_documento("ers.pdf", salida.getvalue())

    assert documento.formato_origen is FormatoDocumento.PDF
    assert documento.contenido_extraido.index("Primera pagina") < (
        documento.contenido_extraido.index("Segunda pagina")
    )


def test_extrae_parrafos_y_tablas_docx_en_orden() -> None:
    documento = cargar_documento("ers.docx", _crear_docx())

    texto = documento.contenido_extraido
    assert documento.formato_origen is FormatoDocumento.DOCX
    assert texto.index("Requisito de transparencia") < texto.index("Responsable")
    assert texto.index("Responsable") < texto.index("Fin del documento")
    assert "Responsable\tEquipo de IA" in texto


def test_carga_desde_ruta(tmp_path: Path) -> None:
    ruta = tmp_path / "ers.txt"
    ruta.write_text("Contenido de la ERS", encoding="utf-8")

    documento = cargar_documento_desde_ruta(ruta)

    assert documento.nombre_original == "ers.txt"
    assert documento.contenido_extraido == "Contenido de la ERS"


def test_ruta_inexistente() -> None:
    with pytest.raises(FileNotFoundError):
        cargar_documento_desde_ruta("no-existe.txt")


@pytest.mark.parametrize("nombre", ["ers.md", "ers.doc", "ers"])
def test_rechaza_formato_no_soportado(nombre: str) -> None:
    with pytest.raises(FormatoNoSoportadoError):
        cargar_documento(nombre, b"contenido")


@pytest.mark.parametrize("contenido", [b"", b"   \n"])
def test_rechaza_txt_vacio(contenido: bytes) -> None:
    with pytest.raises(DocumentoVacioError):
        cargar_documento("ers.txt", contenido)


def test_rechaza_txt_con_codificacion_invalida() -> None:
    with pytest.raises(DocumentoIlegibleError):
        cargar_documento("ers.txt", b"\xff\xfe\xfa")


@pytest.mark.parametrize("nombre", ["ers.pdf", "ers.docx"])
def test_rechaza_documento_corrupto(nombre: str) -> None:
    with pytest.raises(DocumentoIlegibleError):
        cargar_documento(nombre, b"contenido que no corresponde al formato")


def test_rechaza_pdf_sin_texto_extraible() -> None:
    escritor = PdfWriter()
    escritor.add_blank_page(width=612, height=792)
    salida = BytesIO()
    escritor.write(salida)

    with pytest.raises(DocumentoIlegibleError):
        cargar_documento("escaneado.pdf", salida.getvalue())
